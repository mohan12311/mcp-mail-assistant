"""
MCP Mail Assistant - Word 문서 파서
.docx 파일에서 텍스트와 표 추출
"""
import os
from typing import Optional
from core.models import AttachmentAnalysis


def parse_word(filepath: str) -> Optional[AttachmentAnalysis]:
    """Word 문서 파싱"""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx 패키지가 필요합니다: pip install python-docx")
    
    if not os.path.exists(filepath):
        return None
    
    filename = os.path.basename(filepath)
    doc = Document(filepath)
    
    # 텍스트 추출
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    
    extracted_text = "\n".join(paragraphs)
    
    # 표 추출
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        
        if table_data:
            # 첫 행을 헤더로 가정
            if len(table_data) > 1:
                headers = table_data[0]
                rows = table_data[1:]
                tables.append({
                    "headers": headers,
                    "rows": rows
                })
            else:
                tables.append({
                    "headers": [],
                    "rows": table_data
                })
    
    return AttachmentAnalysis(
        filename=filename,
        file_type="word",
        extracted_text=extracted_text,
        tables=tables if tables else None,
        key_points=[],
    )
