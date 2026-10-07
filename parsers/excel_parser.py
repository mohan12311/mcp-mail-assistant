"""
MCP Mail Assistant - Excel 파서
.xlsx 파일에서 시트별 데이터 추출
"""
import os
from typing import Optional
from core.models import AttachmentAnalysis


def parse_excel(filepath: str) -> Optional[AttachmentAnalysis]:
    """Excel 파일 파싱"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("openpyxl 패키지가 필요합니다: pip install openpyxl")
    
    if not os.path.exists(filepath):
        return None
    
    filename = os.path.basename(filepath)
    wb = load_workbook(filepath, data_only=True)
    
    sheet_names = wb.sheetnames
    tables = []
    text_parts = []
    
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        
        # 데이터가 있는 범위 확인
        if ws.max_row == 0 or ws.max_column == 0:
            continue
        
        # 최대 100행, 20열까지만 읽기 (너무 큰 파일 방지)
        max_rows = min(ws.max_row, 100)
        max_cols = min(ws.max_column, 20)
        
        sheet_data = []
        for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols):
            row_data = []
            for cell in row:
                value = cell.value
                if value is not None:
                    row_data.append(str(value))
                else:
                    row_data.append("")
            
            # 빈 행이 아닌 경우만 추가
            if any(cell for cell in row_data):
                sheet_data.append(row_data)
        
        if sheet_data:
            # 첫 행을 헤더로 가정
            headers = sheet_data[0] if sheet_data else []
            rows = sheet_data[1:] if len(sheet_data) > 1 else []
            
            tables.append({
                "sheet_name": sheet_name,
                "headers": headers,
                "rows": rows,
                "total_rows": ws.max_row,
            })
            
            # 텍스트 요약
            text_parts.append(f"[시트: {sheet_name}]")
            text_parts.append(f"컬럼: {', '.join(headers)}")
            text_parts.append(f"데이터 행 수: {len(rows)}")
    
    wb.close()
    
    extracted_text = "\n".join(text_parts)
    
    return AttachmentAnalysis(
        filename=filename,
        file_type="excel",
        extracted_text=extracted_text,
        tables=tables if tables else None,
        sheet_names=sheet_names,
        key_points=[],
    )
