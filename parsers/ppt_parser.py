"""
MCP Mail Assistant - PowerPoint 파서
.pptx 파일에서 슬라이드별 텍스트 추출
"""
import os
from typing import Optional
from core.models import AttachmentAnalysis


def parse_ppt(filepath: str) -> Optional[AttachmentAnalysis]:
    """PowerPoint 파일 파싱"""
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx 패키지가 필요합니다: pip install python-pptx")
    
    if not os.path.exists(filepath):
        return None
    
    filename = os.path.basename(filepath)
    prs = Presentation(filepath)
    
    text_parts = []
    slide_count = len(prs.slides)
    
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        
        # 슬라이드의 모든 shape에서 텍스트 추출
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        
        if slide_texts:
            text_parts.append(f"[슬라이드 {i}]")
            text_parts.extend(slide_texts)
            text_parts.append("")  # 빈 줄로 구분
    
    # 발표자 노트 추출
    notes_parts = []
    for i, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                notes_parts.append(f"[슬라이드 {i} 노트]")
                notes_parts.append(notes)
    
    extracted_text = "\n".join(text_parts)
    if notes_parts:
        extracted_text += "\n\n--- 발표자 노트 ---\n" + "\n".join(notes_parts)
    
    return AttachmentAnalysis(
        filename=filename,
        file_type="ppt",
        extracted_text=extracted_text,
        tables=None,
        page_count=slide_count,
        key_points=[],
    )
